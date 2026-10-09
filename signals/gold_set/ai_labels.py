"""Convert the AI-labelled gold set to labels_AI.csv (v0.1 only; Amendment 7).

Reads the Main tab of the AI workbook named in config/release.yaml (opened read-only; nothing is
written to any workbook, and the owner's human sheet gold_set_L1.xlsx is never opened). Writes
signals/gold_set/labels_AI.csv with: ID, Answer (word), Code (Clarification 19), Unsure (Y or
blank), Note, Labeler. Refuses wrong IDs and any blank or invalid answer, as lock-labels does.

The file must then be locked (make lock-labels FILE=signals/gold_set/labels_AI.csv) before Jev
sees any gold-set passage.

Run: python -m signals.gold_set.ai_labels   (or: make ai-labels)
"""

import csv
from pathlib import Path

import yaml
from openpyxl import load_workbook

from signals.gold_set.labels import WORD_TO_CODE, Refused, expected_ids, to_word

ROOT = Path(__file__).resolve().parents[2]
FIELDS = ["ID", "Answer", "Code", "Unsure", "Note", "Labeler"]


def release():
    with open(ROOT / "config" / "release.yaml") as f:
        return yaml.safe_load(f)


def read_main_tab(path):
    """Rows of the Main tab as dicts, read-only."""
    wb = load_workbook(path, read_only=True)
    rows = list(wb["Main"].iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() for h in rows[0]]
    if header[:5] != ["ID", "Passage", "Answer", "Unsure", "Note"]:
        raise Refused("the Main tab needs the columns ID, Passage, Answer, Unsure, Note")
    return [dict(zip(header, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]


def convert(rows, labeler):
    """Validated output rows; Refused names problem rows by ID only, never an answer."""
    ids = [str(r["ID"]).strip() for r in rows]
    want = expected_ids("main")
    if len(ids) != len(set(ids)) or set(ids) != want:
        raise Refused(f"the IDs do not match the main round exactly ({len(set(ids) & want)} of {len(want)} present)")
    out, bad = [], []
    for r, i in zip(rows, ids):
        word = to_word(str(r["Answer"] or ""))
        unsure = str(r["Unsure"] or "").strip().upper()
        if word is None or unsure not in ("", "Y"):
            bad.append(i)
            continue
        out.append({"ID": i, "Answer": word, "Code": WORD_TO_CODE[word], "Unsure": unsure,
                    "Note": str(r["Note"] or "").strip(), "Labeler": labeler})
    if bad:
        raise Refused(f"{len(bad)} row(s) with a blank or invalid Answer or Unsure: {', '.join(sorted(bad)[:10])}")
    return sorted(out, key=lambda r: r["ID"])


def run():
    cfg = release()
    if cfg.get("release") != "v0.1":
        raise SystemExit("AI labels are used only in release v0.1 (Amendment 7); config/release.yaml says otherwise")
    a = cfg["ai_labels"]
    rows = convert(read_main_tab(ROOT / a["source"]), a["labeler"])
    with open(ROOT / a["output"], "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {a['output']}: {len(rows)} rows, labeler {a['labeler']}")


if __name__ == "__main__":
    run()
