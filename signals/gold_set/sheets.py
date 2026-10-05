"""Build the two labelers' Excel files (session M2.3; Clarifications 19 and 24).

Each file has a Practice tab and a Main tab, in the same order for both labelers. Each row shows
only a neutral ID (PR01..., G001...) and the passage text: no source, date, link, period or type
(Clarification 19). The labeler fills three shaded columns: Answer (a dropdown of five words,
Clarification 26), Unsure (dropdown Y) and Note. The files' author, company and title properties are left blank.

Output: signals/gold_set/sheets/gold_set_L1.xlsx and gold_set_L2.xlsx (git-ignored).
Run: python -m signals.gold_set.sheets   (or: make gold-sheets)
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from signals.corpus.build_corpus import PRIVATE_PATH
from signals.gold_set.draw import KEY_PATH
from signals.gold_set.labels import ANSWER_WORDS, WORD_TO_CODE

ANSWER_LIST = '"' + ",".join(ANSWER_WORDS) + '"'   # Excel list validation, in the order shown

HERE = Path(__file__).resolve().parent
SHEET_DIR = HERE / "sheets"
COLUMNS = ["ID", "Passage", "Answer", "Unsure", "Note"]
FILL_IN = PatternFill("solid", fgColor="DDEBF7")     # light blue: the columns the labeler fills


def load_rounds():
    """{'practice': [(gold_id, text)], 'main': [...]} in key order (the drawn order)."""
    with open(PRIVATE_PATH, newline="") as f:
        text = {r["id"]: r["passage"] for r in csv.DictReader(f)}
    rounds = {"practice": [], "main": []}
    with open(KEY_PATH, newline="") as f:
        for k in csv.DictReader(f):
            rounds[k["round"]].append((k["gold_id"], text[k["corpus_id"]]))
    return rounds


def _tab(ws, items):
    ws.append(COLUMNS)
    for gold_id, passage in items:
        ws.append([gold_id, passage, None, None, None])
    n = len(items) + 1
    answer = _answer_validation()
    unsure = DataValidation(type="list", formula1='"Y"', allow_blank=True, showErrorMessage=True,
                            errorTitle="Unsure", error="Enter Y, or leave blank.")
    for dv, col in ((answer, "C"), (unsure, "D")):
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    for width, col in zip([8, 100, 16, 9, 30], "ABCDE"):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=1, max_row=n):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row > 1 and cell.column_letter in "CDE":
                cell.fill = FILL_IN
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _answer_validation():
    return DataValidation(type="list", formula1=ANSWER_LIST, allow_blank=True, showErrorMessage=True,
                          errorTitle="Answer", error="Pick one answer from the dropdown.")


def update_answer_dropdown(path):
    """Clarification 26, for a sheet already handed out: swap the Answer dropdown to the words in place,
    and convert any answers already typed as 1-4 or N to the matching word. Returns the number of cells
    converted (answers themselves are never shown)."""
    from openpyxl import load_workbook
    code_to_word = {c: w for w, c in WORD_TO_CODE.items()}
    wb = load_workbook(path)
    converted = 0
    for ws in wb:
        n = ws.max_row
        # drop the old Answer dropdown (the one on column C); keep the Unsure dropdown on column D
        ws.data_validations.dataValidation = [dv for dv in ws.data_validations.dataValidation
                                              if not str(dv.sqref).startswith("C")]
        dv = _answer_validation()
        ws.add_data_validation(dv)
        dv.add(f"C2:C{n}")
        ws.column_dimensions["C"].width = 16           # room for "Clear Distress" and "Not Applicable"
        for r in range(2, n + 1):
            cell = ws.cell(r, 3)
            v = str(cell.value).strip().upper() if cell.value is not None else ""
            if v in code_to_word:
                cell.value = code_to_word[v]
                converted += 1
    p = wb.properties
    p.creator = p.lastModifiedBy = p.title = p.subject = p.description = p.keywords = p.category = ""
    wb.save(path)
    return converted


def build(path, rounds):
    wb = Workbook()
    ws = wb.active
    ws.title = "Practice"
    _tab(ws, rounds["practice"])
    _tab(wb.create_sheet("Main"), rounds["main"])
    p = wb.properties                     # nothing that identifies who made the file
    p.creator = p.lastModifiedBy = p.title = p.subject = p.description = p.keywords = p.category = ""
    wb.save(path)


def run(labelers=("L1", "L2"), overwrite=False):
    """Build sheets that do not exist yet. A sheet already handed out may hold a labeler's answers, so it
    is rebuilt only when named and overwrite=True (python -m signals.gold_set.sheets --overwrite L2)."""
    SHEET_DIR.mkdir(exist_ok=True)
    rounds = load_rounds()
    for labeler in labelers:
        path = SHEET_DIR / f"gold_set_{labeler}.xlsx"
        if path.exists() and not overwrite:
            print(f"kept {path.name}: it already exists and may hold answers (use --overwrite {labeler} to rebuild)")
            continue
        build(path, rounds)
        print(f"wrote {path.name} (practice {len(rounds['practice'])}, main {len(rounds['main'])})")


if __name__ == "__main__":
    import sys
    if len(sys.argv) == 3 and sys.argv[1] == "--overwrite" and sys.argv[2] in ("L1", "L2"):
        run((sys.argv[2],), overwrite=True)
    else:
        run()
