"""Checks on the cost model (session M1.8).

Mechanics only: A costs zero by construction, each component matches the
worked-example spreadsheet and a hand calculation, banks that don't take part
cost nothing, the spread sweep moves only the components it should, and the
same seed gives the same table. Nothing here ranks or compares policies.
"""

import numpy as np
import pytest
from openpyxl import load_workbook
from pycel import ExcelCompiler

from engine.cost_workbook import EXAMPLE_POLICIES, STEPS, example_banks, write_cost_workbook
from engine.costs import load_cost_settings, policy_costs
from engine.policies import load_policies
from engine.write_costs import build_cost_table, build_setups
from engine.write_policy_table import POLICY_ORDER

POL = load_policies()
COST = load_cost_settings()
BANKS, SETUPS = build_setups(pol_cfg=POL)
COSTS = {name: policy_costs(SETUPS[name], SETUPS["A"], COST, POL) for name in POLICY_ORDER}
COMPONENTS = ["prepositioning_m", "draws_m", "extra_reserves_m", "hqla_released_m", "annual_cost_m", "one_time_loss_m"]


def test_a_costs_zero():
    for k, v in COSTS["A"].items():
        if k.endswith("_m") or k.endswith("_bn"):
            assert np.all(v == 0), k


def test_one_test_draw_by_hand():
    # $50m overnight at 10 bp a year, actual/360: 50 x 0.0010 / 360 = $0.000139m; four a year.
    for name in ("B", "B_prime", "E"):
        np.testing.assert_allclose(COSTS[name]["draws_m"], 4 * 50 * 0.0010 / 360)


def test_b_and_e_prepositioning_identical():
    np.testing.assert_array_equal(COSTS["B"]["prepositioning_m"], COSTS["E"]["prepositioning_m"])
    assert np.all(COSTS["E"]["extra_reserves_m"] == 0)


@pytest.mark.parametrize("name", ["C", "C_prime"])
def test_non_opt_in_banks_cost_nothing(name):
    out = ~SETUPS[name]["opted_in"]
    assert out.any()
    for k in COMPONENTS:
        assert np.all(COSTS[name][k][out] == 0), k


def test_c_prime_has_no_draw_cost():
    assert np.all(COSTS["C_prime"]["draws_m"] == 0)
    opt = SETUPS["C"]["opted_in"]
    np.testing.assert_allclose(COSTS["C"]["draws_m"][opt],
                               10 * SETUPS["C"]["usage_draw_bn"][opt] * 1000 * 0.0010 / 360)


@pytest.mark.parametrize("name", POLICY_ORDER)
def test_spread_sweep_scales_the_right_components(name):
    lo = policy_costs(SETUPS[name], SETUPS["A"], COST, POL, spread_bp=200)
    hi = policy_costs(SETUPS[name], SETUPS["A"], COST, POL, spread_bp=300)
    mid = COSTS[name]
    for k in ("extra_reserves_m", "hqla_released_m"):           # move with the spread
        np.testing.assert_allclose(lo[k], 0.8 * mid[k])
        np.testing.assert_allclose(hi[k], 1.2 * mid[k])
    for k in ("prepositioning_m", "draws_m", "one_time_loss_m"):  # do not
        np.testing.assert_array_equal(lo[k], mid[k])
        np.testing.assert_array_equal(hi[k], mid[k])


def test_annual_cost_is_the_sum_and_excludes_one_time_loss():
    for name in POLICY_ORDER:
        c = COSTS[name]
        np.testing.assert_allclose(c["annual_cost_m"],
                                   c["prepositioning_m"] + c["draws_m"] + c["extra_reserves_m"] + c["hqla_released_m"])
    # A bank that releases securities has a one-time loss, kept apart (Clarification 13).
    assert (COSTS["C"]["one_time_loss_m"] > 0).any()


def test_same_seed_same_table():
    t1, t2 = build_cost_table(), build_cost_table()
    assert all(np.array_equal(t1[k], t2[k]) for k in t1)
    assert len(t1["policy"]) == 240


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    return write_cost_workbook(BANKS, SETUPS, COST, POL, tmp_path_factory.mktemp("costs") / "costs.xlsx")


# pycel uses Python features due for removal in 3.14; harmless on 3.13, so muted here.
@pytest.mark.filterwarnings("ignore::DeprecationWarning")
def test_workbook_matches_code(workbook):
    sheets = load_workbook(workbook)
    excel = ExcelCompiler(filename=str(workbook))
    rows = example_banks(BANKS, SETUPS["C"])
    assert [BANKS["archetype"][i] for i in rows] == ["svb_like", "gsib"]
    steps = {k for k, *_ in STEPS}
    checked = 0
    for i in rows:
        for name in EXAMPLE_POLICIES:
            title = f"{BANKS['bank_id'][i]} {name}"
            ws = sheets[title]
            for r in range(1, ws.max_row + 1):
                key = ws.cell(r, 5).value
                if key in steps:
                    assert str(ws.cell(r, 2).value).startswith("="), f"{title} {key} is not a formula"
                    value = excel.evaluate(f"'{title}'!B{r}")
                    assert value == pytest.approx(COSTS[name][key][i], abs=1e-9), f"{title} {key}"
                    checked += 1
    assert checked == len(rows) * len(EXAMPLE_POLICIES) * len(STEPS)
