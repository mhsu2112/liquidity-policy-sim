"""Checks on the gold-set draw, labeling sheets and label tools (session M2.3; Clarifications 19 and 24).

Mechanics only: the draw keeps the caps and rules and repeats exactly; sheets show no source,
date, link or period and carry no author; the lock refuses blank, invalid or mismatched files and
never prints answers; compare refuses main-round files; publish needs both locks and matching
fingerprints. No labels are judged, and nothing says any policy or passage is better.
"""

import csv
import random
import re
from collections import Counter

import pytest
from openpyxl import load_workbook

import signals.gold_set.labels as labels
import signals.gold_set.sheets as sheets
from signals.corpus.build_corpus import CORPUS_PATH
from signals.gold_set.draw import PRACTICE, SEED, TYPE_CAP, draw

STRATA = ["S-A", "S-B", "S-C", "S-D"]


@pytest.fixture(scope="module")
def drawn():
    if not CORPUS_PATH.exists():
        pytest.skip("corpus.csv not present")
    with open(CORPUS_PATH, newline="") as f:
        rows = [r for r in csv.DictReader(f) if r["gold_set_eligible"] == "Y"]
    return rows, draw(rows, STRATA, random.Random(SEED))


def test_draw_follows_clarification_24(drawn):
    rows, (main, practice) = drawn
    assert len(main) == 278 and len(practice) == PRACTICE
    assert max(Counter(r["source_type"] for r in main).values()) <= TYPE_CAP
    docs = [r["document_id"] for r in main + practice]
    assert len(docs) == len(set(docs))                                  # one passage per document, rounds disjoint
    sa_docs = {r["document_id"] for r in rows if r["stratum"] == "S-A"}
    assert {r["document_id"] for r in main if r["stratum"] == "S-A"} == sa_docs   # 2007-09: every document
    assert sorted(Counter(r["stratum"] for r in main if r["stratum"] != "S-A").values()) == [74, 74, 75]
    assert all(r["gold_set_eligible"] == "Y" for r in main + practice)


def test_draw_repeats_exactly(drawn):
    rows, (main, practice) = drawn
    again = draw(rows, STRATA, random.Random(SEED))
    assert [r["id"] for r in again[0]] == [r["id"] for r in main]
    assert [r["id"] for r in again[1]] == [r["id"] for r in practice]


def test_sheets_show_only_ids_and_text_and_carry_no_author(tmp_path):
    rounds = {"practice": [("PR01", "Banks borrowed from the discount window.")],
              "main": [("G001", "Lenders tapped the facility."), ("G002", "The Bank turned to the Fed.")]}
    path = tmp_path / "s.xlsx"
    sheets.build(path, rounds)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Practice", "Main"]
    for ws in wb:
        assert [c.value for c in ws[1]] == ["ID", "Passage", "Answer", "Unsure", "Note"]
        text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
        assert not re.search(r"https?://|www\.|\b(19|20)\d\d-\d\d-\d\d\b|\bS-[ABCD]\b|filing|news|speech|official", text)
        formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        assert '"Reassuring,Routine,Some Concern,Clear Distress,Not Applicable"' in formulas and '"Y"' in formulas
    p = wb.properties
    assert not p.creator and not p.lastModifiedBy and not p.title


def _setup(tmp_path, monkeypatch, n_main=3, n_practice=2):
    summary = tmp_path / "summary.csv"
    summary.write_text(f"round,stratum,source_type,count\nmain,S-B,news,{n_main}\npractice,S-B,news,{n_practice}\n")
    monkeypatch.setattr(labels, "SUMMARY_PATH", summary)
    monkeypatch.setattr(labels, "INCOMING", tmp_path / "incoming")
    monkeypatch.setattr(labels, "LOCKS_PATH", tmp_path / "locks.md")
    calls = []
    monkeypatch.setattr(labels, "git", lambda *a: calls.append(a))
    return calls


def _write(path, rows):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["ID", "Passage", "Answer", "Unsure", "Note"])
        w.writerows(rows)
    return path


def test_lock_refuses_a_file_with_a_blank_answer(tmp_path, monkeypatch):
    calls = _setup(tmp_path, monkeypatch)
    f = _write(tmp_path / "gold_set_L2_main.csv", [["G001", "t", "Routine", "", ""], ["G002", "t", "", "", ""], ["G003", "t", "Clear Distress", "Y", ""]])
    with pytest.raises(labels.Refused, match="1 blank answer"):
        labels.lock(f)
    assert not (tmp_path / "locks.md").exists() and not calls


def test_lock_refuses_invalid_answers_wrong_ids_and_wrong_names(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    bad = _write(tmp_path / "gold_set_L1_main.csv", [["G001", "t", "Worrying", "", ""], ["G002", "t", "Routine", "maybe", ""], ["G003", "t", "Reassuring", "", ""]])
    with pytest.raises(labels.Refused, match="not one of Reassuring"):
        labels.lock(bad)
    numbers = _write(tmp_path / "gold_set_L1_main.csv", [["G001", "t", "1", "", ""], ["G002", "t", "2", "", ""], ["G003", "t", "N", "", ""]])
    with pytest.raises(labels.Refused, match="3 answer"):      # Clarification 26: words only
        labels.lock(numbers)
    short = _write(tmp_path / "gold_set_L1_main.csv", [["G001", "t", "Reassuring", "", ""], ["G002", "t", "Routine", "", ""]])
    with pytest.raises(labels.Refused, match="do not match"):
        labels.lock(short)
    named = _write(tmp_path / "my_labels.csv", [["G001", "t", "Routine", "", ""]])
    with pytest.raises(labels.Refused, match="file name"):
        labels.lock(named)


def test_lock_records_fingerprint_commits_locks_only_and_never_prints_answers(tmp_path, monkeypatch, capsys):
    calls = _setup(tmp_path, monkeypatch)
    f = _write(tmp_path / "gold_set_L1_main.csv", [["G001", "t", "Routine", "", "SECRET-NOTE"], ["G002", "t", "  not   applicable ", "", ""], ["G003", "t", "clear distress", "Y", ""]])
    labels.lock(f)
    out = capsys.readouterr().out
    assert "SECRET-NOTE" not in out and "Routine" not in out and "distress" not in out.lower()
    line = (tmp_path / "locks.md").read_text().strip().splitlines()[-1]
    assert "| L1 | main | gold_set_L1_main.csv | 3 |" in line and labels.sha256(f) in line
    assert calls[1][0] == "commit" and calls[1][-1].endswith("locks.md") and calls[2] == ("push",)
    with pytest.raises(labels.Refused, match="already locked"):
        labels.lock(f)


def test_compare_practice_refuses_main_round_files(tmp_path, monkeypatch):
    _setup(tmp_path, monkeypatch)
    (tmp_path / "incoming").mkdir()
    for l in ("L1", "L2"):   # main-round IDs saved under practice names
        _write(tmp_path / "incoming" / f"gold_set_{l}_practice.csv", [["G001", "t", "Reassuring", "", ""], ["G002", "t", "Routine", "", ""]])
    with pytest.raises(labels.Refused):
        labels.compare_practice()


def test_publish_needs_both_locks_and_matching_fingerprints(tmp_path, monkeypatch):
    calls = _setup(tmp_path, monkeypatch)
    monkeypatch.setattr(labels, "HERE", tmp_path / "signals" / "gold_set")
    rows = [["G001", "t", "Reassuring", "", ""], ["G002", "t", "Routine", "", ""], ["G003", "t", "Some Concern", "", ""]]
    labels.lock(_write(tmp_path / "gold_set_L1_main.csv", rows))
    with pytest.raises(labels.Refused, match="L2 main round is not locked"):
        labels.publish()
    labels.lock(_write(tmp_path / "gold_set_L2_main.csv", rows))
    _write(tmp_path / "incoming" / "gold_set_L2_main.csv", rows[:2] + [["G003", "t", "Clear Distress", "", ""]])   # changed after lock
    with pytest.raises(labels.Refused, match="no longer matches"):
        labels.publish()


def test_publish_writes_only_id_answer_unsure_note_once_both_locks_match(tmp_path, monkeypatch):
    calls = _setup(tmp_path, monkeypatch)
    out_dir = tmp_path / "signals" / "gold_set"
    out_dir.mkdir(parents=True)
    monkeypatch.setattr(labels, "HERE", out_dir)
    for l in ("L1", "L2"):
        labels.lock(_write(tmp_path / f"gold_set_{l}_main.csv",
                           [["G001", "passage text", "reassuring", "", ""], ["G002", "passage text", "Some Concern", "y", "odd"],
                            ["G003", "passage text", "Not Applicable", "", ""]]))
    labels.publish()
    with open(out_dir / "labels_L1.csv") as f:
        rows = list(csv.DictReader(f))
    assert list(rows[0]) == ["ID", "Answer", "Code", "Unsure", "Note"] and [r["ID"] for r in rows] == ["G001", "G002", "G003"]
    assert [(r["Answer"], r["Code"]) for r in rows] == [("Reassuring", "1"), ("Some Concern", "3"), ("Not Applicable", "N")]
    assert rows[1]["Unsure"] == "Y" and "Passage" not in rows[0]
    assert calls[-1][0] == "commit"


def test_word_labels_map_to_the_clarification_19_codes():
    assert labels.WORD_TO_CODE == {"Reassuring": "1", "Routine": "2", "Some Concern": "3", "Clear Distress": "4",
                                   "Not Applicable": "N"}
    assert labels.to_word("  some   CONCERN ") == "Some Concern" and labels.to_word("3") is None


def test_in_place_update_converts_typed_codes_and_keeps_properties_blank(tmp_path):
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    path = tmp_path / "old.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    ws.append(["ID", "Passage", "Answer", "Unsure", "Note"])
    for row in (["G001", "t", 1, None, None], ["G002", "t", "n", "Y", None], ["G003", "t", None, None, None]):
        ws.append(row)
    old = DataValidation(type="list", formula1='"1,2,3,4,N"')
    ws.add_data_validation(old)
    old.add("C2:C4")
    wb.save(path)
    assert sheets.update_answer_dropdown(path) == 2
    wb = load_workbook(path)
    ws = wb["Main"]
    assert [ws.cell(r, 3).value for r in (2, 3, 4)] == ["Reassuring", "Not Applicable", None]
    assert [dv.formula1 for dv in ws.data_validations.dataValidation] == [sheets.ANSWER_LIST]
    assert not wb.properties.creator and not wb.properties.lastModifiedBy


def test_sheet_builder_never_overwrites_an_existing_sheet(tmp_path, monkeypatch):
    monkeypatch.setattr(sheets, "SHEET_DIR", tmp_path)
    monkeypatch.setattr(sheets, "load_rounds", lambda: {"practice": [("PR01", "a")], "main": [("G001", "b")]})
    existing = tmp_path / "gold_set_L1.xlsx"
    existing.write_bytes(b"labeler's work")
    sheets.run()
    assert existing.read_bytes() == b"labeler's work" and (tmp_path / "gold_set_L2.xlsx").exists()
