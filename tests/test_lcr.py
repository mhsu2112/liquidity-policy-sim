"""Checks on the balance sheet and the LCR under current rules (session M1.2).

Mechanics only: the 40% cap, the calibration, the pro-rata loss split, balance
sheets that balance, and the worked-examples spreadsheet agreeing with the code.
Nothing here compares policies.
"""

import numpy as np
import pytest
from openpyxl import load_workbook
from pycel import ExcelCompiler

from engine.balance_sheet import ONE_DOLLAR_BN, check_balances
from engine.banks import generate_banks
from engine.lcr import compute_lcr, load_lcr_settings
from engine.lcr_workbook import BANK_INPUTS, STEPS, example_rows, write_workbook

SETTINGS = load_lcr_settings()
BANKS = generate_banks()
LCR = compute_lcr(BANKS, SETTINGS)


def hand_bank(level2a=40.0, calibration=0.70, loss=0.0):
    """A one-bank balance sheet small enough to calculate by hand ($bn).

    Reserves 10, Level 1 20, Level 2A 40 (by default); insured deposits 100,
    uninsured 50, wholesale funding 10.
    """
    a = lambda x: np.array([float(x)])  # noqa: E731
    return {"bank_id": np.array(["HAND-01"]), "reserves_bn": a(10), "level1_securities_bn": a(20),
            "level2a_securities_bn": a(level2a), "securities_bn": a(20 + level2a),
            "unrealized_loss_bn": a(loss), "loans_bn": a(100), "insured_deposits_bn": a(100),
            "uninsured_deposits_bn": a(50), "stwf_bn": a(10), "equity_bn": a(10),
            "lcr_calibration": a(calibration)}


def test_cap_binds_by_hand():
    r = compute_lcr(hand_bank(), SETTINGS)
    # Level 1 = 10 + 20 = 30. Level 2A after haircut = 40 x 0.85 = 34.
    # Cap = 30 x 40/60 = 20, so only 20 counts and 14 is excluded.
    assert r["level1_hqla_bn"][0] == pytest.approx(30)
    assert r["level2a_after_haircut_bn"][0] == pytest.approx(34)
    assert r["level2a_counted_bn"][0] == pytest.approx(20)
    assert r["level2a_excluded_by_cap_bn"][0] == pytest.approx(14)
    assert r["hqla_bn"][0] == pytest.approx(50)
    assert r["level2a_counted_bn"][0] / r["hqla_bn"][0] == pytest.approx(0.40)


def test_cap_does_not_bind_by_hand():
    r = compute_lcr(hand_bank(level2a=10), SETTINGS)
    # Level 2A after haircut = 8.5, below the cap of 20, so it all counts.
    assert r["level2a_counted_bn"][0] == pytest.approx(8.5)
    assert r["level2a_excluded_by_cap_bn"][0] == pytest.approx(0)
    assert r["hqla_bn"][0] == pytest.approx(38.5)


def test_lcr_by_hand():
    r = compute_lcr(hand_bank(), SETTINGS)
    # Outflows = 3% x 100 + 40% x 50 + 100% x 10 = 3 + 20 + 10 = 33.
    # At 70% calibration: 23.1. LCR = 50 / 23.1 = 216.45%.
    assert r["gross_outflows_bn"][0] == pytest.approx(33)
    assert r["net_cash_outflows_bn"][0] == pytest.approx(33)
    assert r["calibrated_outflows_bn"][0] == pytest.approx(23.1)
    assert r["lcr"][0] == pytest.approx(50 / 23.1)


def test_calibration_scales_outflows():
    full, cat3, svb = (compute_lcr(hand_bank(calibration=c), SETTINGS)["lcr"][0] for c in (1.00, 0.85, 0.70))
    assert cat3 == pytest.approx(full / 0.85)
    assert svb == pytest.approx(full / 0.70)


def test_each_type_gets_its_calibration():
    # Contract section 1c, typed in from the contract text.
    expected = {"gsib": 1.00, "regional_cat3": 0.85, "svb_like": 0.70, "diversified_regional": 1.00}
    for name, factor in expected.items():
        rows = BANKS["archetype"] == name
        assert (LCR["outflow_factor"][rows] == factor).all(), name
        np.testing.assert_allclose(LCR["calibrated_outflows_bn"][rows],
                                   LCR["net_cash_outflows_bn"][rows] * factor)
    no_lcr = BANKS["archetype"] == "diversified_regional"
    assert (LCR["lcr_status"][no_lcr] == "not required").all()
    assert (LCR["lcr_status"][~no_lcr] == "required").all()
    assert np.isfinite(LCR["lcr"][no_lcr]).all()  # ratio still computed


def test_cap_never_exceeded_for_any_bank():
    share = LCR["level2a_counted_bn"] / LCR["hqla_bn"]
    assert (share <= SETTINGS["hqla"]["level2_cap"] + 1e-12).all()


def test_losses_split_pro_rata():
    mv = LCR["level1_securities_mv_bn"] + LCR["level2a_securities_mv_bn"]
    np.testing.assert_allclose(mv, BANKS["securities_bn"] - BANKS["unrealized_loss_bn"], atol=ONE_DOLLAR_BN)
    np.testing.assert_allclose(LCR["level1_securities_mv_bn"] / BANKS["level1_securities_bn"],
                               LCR["level2a_securities_mv_bn"] / BANKS["level2a_securities_bn"])


def test_no_inflows():
    assert (LCR["inflows_bn"] == 0).all()
    np.testing.assert_array_equal(LCR["net_cash_outflows_bn"], LCR["gross_outflows_bn"])


def test_balance_sheets_still_balance():
    check_balances(BANKS)  # raises if any bank is out by $1 or more, book or market value
    np.testing.assert_allclose(LCR["mtm_equity_bn"], BANKS["equity_bn"] - BANKS["unrealized_loss_bn"])


def test_check_balances_catches_an_error():
    # The check itself must fail when a balance sheet is broken.
    broken = {k: v.copy() for k, v in BANKS.items()}
    broken["reserves_bn"][0] += 0.001  # $1 million out
    with pytest.raises(ValueError):
        check_balances(broken)


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    path = tmp_path_factory.mktemp("xlsx") / "lcr_worked_examples.xlsx"
    write_workbook(BANKS, LCR, path, SETTINGS)
    return path


# pycel uses Python features due for removal in 3.14; harmless on 3.13, so muted here.
@pytest.mark.filterwarnings("ignore::DeprecationWarning")
def test_workbook_matches_code(workbook):
    sheets = load_workbook(workbook)
    excel = ExcelCompiler(filename=str(workbook))
    rows = example_rows(BANKS)
    assert [BANKS["bank_id"][i] for i in rows] == sheets.sheetnames[1:]
    bank_inputs = {k for k, _, _ in BANK_INPUTS}
    steps = {k for k, *_ in STEPS if k in LCR}  # calculation rows the code also reports
    for i in rows:
        ws = sheets[BANKS["bank_id"][i]]
        checked = 0
        for r in range(1, ws.max_row + 1):
            key, value = ws.cell(r, 5).value, ws.cell(r, 2).value
            if key in bank_inputs:
                # Inputs are plain numbers equal to the model's balance sheet
                # (Excel keeps 15 significant digits, far finer than $1).
                assert not isinstance(value, str) and value == pytest.approx(BANKS[key][i], abs=ONE_DOLLAR_BN)
            elif key in steps:
                # Calculations are formulas whose answer equals the code's.
                assert isinstance(value, str) and value.startswith("=")
                got = excel.evaluate(f"'{ws.title}'!B{r}")
                assert got == pytest.approx(LCR[key][i], rel=1e-12, abs=ONE_DOLLAR_BN), (ws.title, key)
                checked += 1
        assert checked >= 14, ws.title  # every step, ending in the LCR, was compared
