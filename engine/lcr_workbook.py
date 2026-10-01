"""Worked-examples spreadsheet for the LCR (session M1.2).

Three banks (the first SVB-like, Category III regional and GSIB) are worked
through in Excel with live formulas, so each step can be checked by hand.
The inputs are typed-in numbers taken from the model; every calculation cell is
a formula that refers only to cells on the same sheet. The code never writes
an answer into a calculation cell, so the spreadsheet is an independent check.
tests/test_lcr.py works the formulas out and confirms they match the code.
"""

import re
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from engine.lcr import load_lcr_settings
from engine.write_banks import ROOT, git_commit

DEFAULT_WORKBOOK = ROOT / "outputs" / "lcr_worked_examples.xlsx"
EXAMPLE_TYPES = ["svb_like", "regional_cat3", "gsib"]  # session M1.2 brief

MONEY, RATE = "#,##0.000", "0.00%"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # pale yellow: typed-in inputs
BOLD = Font(bold=True)

# Inputs: (key, label, where the number comes from). Bank lines come from the
# model's balance sheet; rules come from config/lcr.yaml.
BANK_INPUTS = [
    ("reserves_bn", "Reserves (cash at the Fed)", "banks.csv"),
    ("level1_securities_bn", "Level 1 securities, book value", "banks.csv"),
    ("level2a_securities_bn", "Level 2A securities, book value", "banks.csv"),
    ("unrealized_loss_bn", "Unrealized loss on securities", "banks.csv"),
    ("insured_deposits_bn", "Insured deposits", "banks.csv"),
    ("uninsured_deposits_bn", "Uninsured deposits", "banks.csv"),
    ("stwf_bn", "Short-term wholesale funding", "banks.csv"),
    ("loans_bn", "Loans (base for inflows)", "banks.csv"),
    ("equity_bn", "Equity, book value", "banks.csv"),
]
RULE_INPUTS = [
    ("level1_haircut", "Level 1 haircut", "contract 7; 12 CFR 249.21"),
    ("level2a_haircut", "Level 2A haircut", "contract 7; 12 CFR 249.21"),
    ("level2_cap", "Level 2 cap, share of total HQLA", "contract 7; 12 CFR 249.21"),
    ("rate_insured", "Outflow rate: insured deposits", "Clarification 3; 249.32(a)(1)"),
    ("rate_uninsured", "Outflow rate: uninsured deposits", "Clarification 3; 249.32(h)(1)(ii)"),
    ("rate_stwf", "Outflow rate: short-term wholesale funding", "Clarification 3; 249.32(h)(3)"),
    ("loan_inflow_rate", "Inflow rate on loans", "Clarification 3: none modeled"),
    ("inflow_cap", "Inflow cap, share of outflows", "12 CFR 249.30"),
    ("outflow_factor", "Calibration (outflow factor)", "contract 1c"),
]
# Calculation: (key, label, formula with {key} placeholders, note, format).
# Keys ending in _bn or equal to "lcr" match columns in outputs/banks_lcr.csv.
STEPS = [
    ("loss_rate", "Loss rate on securities", "={unrealized_loss_bn}/({level1_securities_bn}+{level2a_securities_bn})",
     "Unrealized loss as a share of all securities; applied pro rata", RATE),
    ("level1_securities_mv_bn", "Level 1 securities, market value", "={level1_securities_bn}*(1-{loss_rate})",
     "Book value less its pro-rata share of the loss", MONEY),
    ("level2a_securities_mv_bn", "Level 2A securities, market value", "={level2a_securities_bn}*(1-{loss_rate})",
     "Book value less its pro-rata share of the loss", MONEY),
    ("level1_hqla_bn", "Level 1 HQLA", "=({reserves_bn}+{level1_securities_mv_bn})*(1-{level1_haircut})",
     "Reserves plus Level 1 securities, no haircut", MONEY),
    ("level2a_after_haircut_bn", "Level 2A after haircut", "={level2a_securities_mv_bn}*(1-{level2a_haircut})",
     "Market value less the 15% haircut", MONEY),
    ("level2a_cap_bn", "Most Level 2A that may count", "={level1_hqla_bn}*{level2_cap}/(1-{level2_cap})",
     "40% of total HQLA = 40/60 of Level 1", MONEY),
    ("level2a_counted_bn", "Level 2A counted", "=MIN({level2a_after_haircut_bn},{level2a_cap_bn})",
     "The smaller of the two lines above", MONEY),
    ("level2a_excluded_by_cap_bn", "Level 2A excluded by the cap", "={level2a_after_haircut_bn}-{level2a_counted_bn}",
     "Zero when the cap does not bind", MONEY),
    ("hqla_bn", "Total HQLA", "={level1_hqla_bn}+{level2a_counted_bn}", "Level 1 plus Level 2A counted", MONEY),
    ("level2a_share_of_hqla", "Level 2A share of HQLA", "={level2a_counted_bn}/{hqla_bn}",
     "Never above 40%", RATE),
    ("gross_outflows_bn", "30-day outflows",
     "={rate_insured}*{insured_deposits_bn}+{rate_uninsured}*{uninsured_deposits_bn}+{rate_stwf}*{stwf_bn}",
     "Each funding line times its run-off rate", MONEY),
    ("inflows_bn", "30-day inflows", "={loan_inflow_rate}*{loans_bn}", "None modeled (Clarification 3)", MONEY),
    ("net_cash_outflows_bn", "Net cash outflows", "={gross_outflows_bn}-MIN({inflows_bn},{inflow_cap}*{gross_outflows_bn})",
     "Outflows less inflows, inflows capped at 75% of outflows", MONEY),
    ("calibrated_outflows_bn", "Calibrated net cash outflows", "={net_cash_outflows_bn}*{outflow_factor}",
     "Reduced LCRs scale outflows down (85% or 70%)", MONEY),
    ("lcr", "LCR", "={hqla_bn}/{calibrated_outflows_bn}", "Total HQLA divided by calibrated outflows", RATE),
    ("mtm_equity_bn", "Mark-to-market equity (information only)", "={equity_bn}-{unrealized_loss_bn}",
     "Book equity less the unrealized loss", MONEY),
]


def example_rows(banks):
    """Index of the first bank of each example type."""
    return [int(np.flatnonzero(banks["archetype"] == t)[0]) for t in EXAMPLE_TYPES]


def _rule_values(s, factor):
    h, r, i = s["hqla"], s["outflow_rates"], s["inflows"]
    return {"level1_haircut": h["level1_haircut"], "level2a_haircut": h["level2a_haircut"],
            "level2_cap": h["level2_cap"], "rate_insured": r["insured_deposits"],
            "rate_uninsured": r["uninsured_deposits"], "rate_stwf": r["stwf"],
            "loan_inflow_rate": i["loan_inflow_rate"], "inflow_cap": i["cap_share_of_outflows"],
            "outflow_factor": factor}


def _bank_sheet(wb, banks, lcr, i, s, stamp):
    ws = wb.create_sheet(str(banks["bank_id"][i]))
    ws["A1"] = f"{banks['bank_id'][i]}: LCR worked example ({banks['archetype'][i]}, Category {banks['lcr_category'][i]})"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = stamp
    for col, width in zip("ABCDE", (44, 16, 8, 52, 30)):
        ws.column_dimensions[col].width = width

    cell = {}  # key -> cell address, so formulas refer to the right row
    row = 4
    for title, items, values in [
        ("Inputs: balance sheet ($bn, typed in from the model)", BANK_INPUTS,
         {k: float(banks[k][i]) for k, _, _ in BANK_INPUTS}),
        ("Inputs: rules (typed in from config/lcr.yaml)", RULE_INPUTS,
         _rule_values(s, float(lcr["outflow_factor"][i]))),
    ]:
        ws.cell(row, 1, title).font = BOLD
        row += 1
        for key, label, source in items:
            ws.cell(row, 1, label)
            c = ws.cell(row, 2, values[key])
            c.fill = INPUT_FILL
            c.number_format = MONEY if key.endswith("_bn") else RATE
            ws.cell(row, 4, source)
            ws.cell(row, 5, key)
            cell[key] = f"B{row}"
            row += 1
        row += 1

    ws.cell(row, 1, "Calculation (live formulas: click a cell to see it)").font = BOLD
    ws.cell(row, 5, "column in banks_lcr.csv").font = BOLD
    row += 1
    for key, label, formula, note, fmt in STEPS:
        cell[key] = f"B{row}"
        ws.cell(row, 1, label)
        c = ws.cell(row, 2, re.sub(r"\{(\w+)\}", lambda m: cell[m.group(1)], formula))
        c.number_format = fmt
        ws.cell(row, 4, note)
        ws.cell(row, 5, key)
        if key == "lcr":
            ws.cell(row, 1).font = c.font = BOLD
        row += 1
    ws.cell(row + 1, 1, "Status").font = BOLD
    ws.cell(row + 1, 2, f"{lcr['lcr_status'][i]} at {lcr['outflow_factor'][i]:.0%} calibration")
    return cell


def write_workbook(banks, lcr, out_path=DEFAULT_WORKBOOK, s=None):
    s = s or load_lcr_settings()
    stamp = f"git commit {git_commit()} | model LCR for comparison in outputs/banks_lcr.csv"
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    readme.column_dimensions["A"].width = 110
    for r, text in enumerate([
        "LCR worked examples (session M1.2)",
        "",
        "One sheet per bank: the first SVB-like, Category III regional and GSIB bank in outputs/banks.csv.",
        "Yellow cells are inputs typed in from the model. Every other number is a live Excel formula.",
        "Column E names the matching column in outputs/banks_lcr.csv, so each step can be compared.",
        "Rules: contract section 7 (haircuts, 40% cap) and docs/amendments.md Clarification 3 (rates).",
        "No discount window credit or other policy rule is included; that comes in session M1.7.",
        "An automated test (tests/test_lcr.py) works out these formulas and checks they match the code.",
        "",
        stamp,
    ], start=1):
        readme.cell(r, 1, text).alignment = Alignment(wrap_text=True)
    readme["A1"].font = Font(bold=True, size=13)

    for i in example_rows(banks):
        _bank_sheet(wb, banks, lcr, i, s, stamp)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
