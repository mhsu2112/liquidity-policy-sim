"""Worked-example spreadsheet for the cost model (session M1.8).

Two banks (the first SVB-like bank and the first GSIB that opt in under C, so
every component appears) are worked through under B, B' and C, one sheet each.
Inputs are typed-in numbers from the model; every cost cell is a live formula
referring only to cells on its own sheet. The code never writes an answer into a
cost cell, so the spreadsheet is an independent check. tests/test_costs.py works
the formulas out and confirms they match engine/costs.py.
"""

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from engine.collateral import LOAN_TYPES, SECURITY_CLASSES
from engine.costs import at_fed, draws_per_year
from engine.write_banks import ROOT, git_commit

DEFAULT_WORKBOOK = ROOT / "outputs" / "cost_worked_example.xlsx"
EXAMPLE_POLICIES = ["B", "B_prime", "C"]   # session M1.8 brief
EXAMPLE_TYPES = ["svb_like", "gsib"]
KINDS = (*SECURITY_CLASSES, *LOAN_TYPES)
LABEL = {"level1": "Level 1 securities (market value)", "level2a": "Level 2A securities (market value)",
         "resi": "Residential loans", "cre": "CRE loans", "ci": "C&I loans"}

MONEY, BPS, NUM = "#,##0.000", "0", "#,##0.000000"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # pale yellow: typed-in inputs
BOLD = Font(bold=True)


def example_banks(banks, setup_c):
    """Index of the first bank of each example type that opts in under C."""
    return [int(next(i for i, (a, o) in enumerate(zip(banks["archetype"], setup_c["opted_in"])) if a == t and o))
            for t in EXAMPLE_TYPES]


def _inputs(setup, setup_a, i, cost_cfg, pol_cfg):
    """(key, label, value, source, format) for one bank under one policy."""
    mine, base = at_fed(setup["placement"]), at_fed(setup_a["placement"])
    count, size = draws_per_year(setup, pol_cfg)
    rows = []
    for k in KINDS:
        rows.append((f"policy_{k}", f"At the Fed under the policy: {LABEL[k]}", float(mine[k][i]), "policy_table.csv", MONEY))
        rows.append((f"a_{k}", f"At the Fed under A: {LABEL[k]}", float(base[k][i]), "policy_table.csv", MONEY))
        rows.append((f"bp_{k}", f"Prepositioning rate: {LABEL[k]} (bp a year)", cost_cfg["prepositioning_bp"][k],
                     "contract 6; config/costs.yaml", BPS))
    rows += [
        ("multiple", "Prepositioning rate multiple", cost_cfg["prepositioning_rate_multiple"], "contract 5 (x1 here)", NUM),
        ("draws", "Routine draws a year", float(count[i]), "tests (B, B', E) or usage draws (C)", NUM),
        ("draw_size", "Size of each draw ($bn)", float(size[i]), "Decision 2-3; Clarification 11 item 7", NUM),
        ("draw_bp", "Primary credit rate - interest on reserves (bp)", cost_cfg["draw_spread_bp"],
         "Fed implementation note, 2026-09-16", BPS),
        ("days", "Day count (days a year)", cost_cfg["day_count_days"], "Clarification 13: actual/360", BPS),
        ("extra_reserves", "Extra reserves ($bn)", float(setup["extra_reserves_bn"][i]), "policy_table.csv", MONEY),
        ("released", "HQLA released into loans ($bn)", float(sum(setup["released"].values())[i]), "policy_table.csv", MONEY),
        ("spread_bp", "Loan-to-reserve spread (bp a year)", cost_cfg["loan_to_reserve_spread_bp"], "contract 6", BPS),
        ("equity_a", "Book equity under A ($bn)", float(setup_a["banks"]["equity_bn"][i]), "policy_table.csv", MONEY),
        ("equity_policy", "Book equity under the policy ($bn)", float(setup["banks"]["equity_bn"][i]),
         "policy_table.csv", MONEY),
    ]
    return rows


# Calculation: (key, label, formula with {key} placeholders, note). Keys match columns in outputs/costs.csv.
STEPS = [
    *[(f"prepositioning_{k}_m", f"Prepositioning: {LABEL[k]} ($m)",
       f"=({{policy_{k}}}-{{a_{k}}})*{{bp_{k}}}/10000*{{multiple}}*1000",
       "(policy - A) x rate; negative if less than under A") for k in KINDS],
    ("prepositioning_m", "Prepositioning, total ($m a year)",
     "=" + "+".join(f"{{prepositioning_{k}_m}}" for k in KINDS), "Sum of the five lines above"),
    ("draws_m", "Routine draws ($m a year)", "={draws}*{draw_size}*{draw_bp}/10000/{days}*1000",
     "Draws x size x spread x one night"),
    ("extra_reserves_m", "Extra reserves ($m a year)", "={extra_reserves}*{spread_bp}/10000*1000",
     "B and B' only"),
    ("hqla_released_m", "HQLA released ($m a year, negative = saving)", "=-{released}*{spread_bp}/10000*1000",
     "C and C' banks that opt in earn the spread"),
    ("annual_cost_m", "Annual cost relative to A ($m)",
     "={prepositioning_m}+{draws_m}+{extra_reserves_m}+{hqla_released_m}", "Sum of the four components"),
    ("one_time_loss_m", "One-time loss on released securities ($m; not in the annual cost)",
     "=({equity_a}-{equity_policy})*1000", "Unrealized loss realized on release"),
]


def _sheet(wb, title, heading, rows, stamp):
    ws = wb.create_sheet(title)
    ws["A1"], ws["A2"] = heading, stamp
    ws["A1"].font = Font(bold=True, size=13)
    for col, width in zip("ABCDE", (62, 16, 4, 40, 26)):
        ws.column_dimensions[col].width = width
    cell, r = {}, 4
    ws.cell(r, 1, "Inputs (typed in from the model and config)").font = BOLD
    r += 1
    for key, label, value, source, fmt in rows:
        ws.cell(r, 1, label)
        c = ws.cell(r, 2, value)
        c.fill, c.number_format = INPUT_FILL, fmt
        ws.cell(r, 4, source)
        ws.cell(r, 5, key)
        cell[key] = f"B{r}"
        r += 1
    r += 1
    ws.cell(r, 1, "Calculation (live formulas: click a cell to see it)").font = BOLD
    ws.cell(r, 5, "column in costs.csv").font = BOLD
    r += 1
    for key, label, formula, note in STEPS:
        cell[key] = f"B{r}"
        ws.cell(r, 1, label)
        c = ws.cell(r, 2, re.sub(r"\{(\w+)\}", lambda m: cell[m.group(1)], formula))
        c.number_format = MONEY
        ws.cell(r, 4, note)
        ws.cell(r, 5, key)
        if key == "annual_cost_m":
            ws.cell(r, 1).font = c.font = BOLD
        r += 1


def write_cost_workbook(banks, setups, cost_cfg, pol_cfg, out_path=DEFAULT_WORKBOOK):
    stamp = f"git commit {git_commit()} | model figures for comparison in outputs/costs.csv"
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    readme.column_dimensions["A"].width = 110
    for r, text in enumerate([
        "Cost worked examples (session M1.8)",
        "",
        "One sheet per bank and policy: the first SVB-like bank and the first GSIB that opt in under C, under B, B' and C.",
        "Yellow cells are inputs typed in from the model and config/costs.yaml. Every other number is a live formula.",
        "Every cost is annual, in $ millions, relative to policy A (docs/amendments.md Clarification 13).",
        "Column E names the matching column in outputs/costs.csv. tests/test_costs.py checks the formulas match the code.",
        "This file compares nothing across policies; costs are reported alongside outcomes in M3 (contract section 4).",
        "",
        stamp,
    ], start=1):
        readme.cell(r, 1, text).alignment = Alignment(wrap_text=True)
    readme["A1"].font = Font(bold=True, size=13)
    for i in example_banks(banks, setups["C"]):
        for name in EXAMPLE_POLICIES:
            bank = str(banks["bank_id"][i])
            _sheet(wb, f"{bank} {name}", f"{bank} under {name}: annual cost relative to A",
                   _inputs(setups[name], setups["A"], i, cost_cfg, pol_cfg), stamp)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
